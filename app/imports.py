from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Ledger, Transaction


MAX_IMPORT_FILE_BYTES = 12 * 1024 * 1024
MAX_IMPORT_ROWS = 5_000
IMPORT_SOURCES = {"wechat-import", "alipay-import"}
FAILED_STATUS_MARKERS = ("失败", "关闭", "撤销", "取消")

ALIPAY_CATEGORY_RULES = (
    ("餐饮", ("餐饮", "美食", "外卖", "咖啡", "茶饮")),
    ("交通", ("交通", "出行", "打车", "停车", "加油", "充电")),
    ("住房", ("住房", "物业", "房租", "家居")),
    ("购物", ("百货", "购物", "数码", "服饰", "商超", "盒马")),
    ("医疗", ("医疗", "健康", "药")),
    ("教育", ("教育", "培训", "书")),
    ("娱乐", ("休闲", "娱乐", "旅游", "游戏")),
    ("转账", ("亲友", "转账", "红包")),
)

GENERAL_CATEGORY_RULES = (
    ("交通", ("停车", "车场", "打车", "地铁", "公交", "加油", "充电", "出行")),
    ("餐饮", ("餐饮", "美食", "餐厅", "外卖", "咖啡", "茶", "冒菜", "盒饭")),
    ("购物", ("盒马", "超市", "百货", "购物", "采购", "商场", "商品")),
    ("住房", ("房租", "住房", "物业", "家居")),
    ("医疗", ("医院", "医疗", "药房", "健康")),
    ("教育", ("教育", "培训", "书店")),
    ("娱乐", ("旅游", "景区", "游乐", "游戏", "酒吧")),
    ("转账", ("转账", "亲友代付", "二维码付款", "红包")),
)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\ufeff", "").replace("\t", "").strip()


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "交易时间" in text and "收/支" in text:
            return text
    raise ValueError("无法识别 CSV 编码，请使用支付宝官方导出的 CSV 文件")


def _xlsx_rows(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Excel 文件无法读取，请重新从微信支付导出") from exc
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(_decode_csv(content))))


def _find_header(rows: list[list[Any]]) -> tuple[int, list[str], str]:
    for index, row in enumerate(rows):
        headers = [_clean(value) for value in row]
        header_set = set(headers)
        if {"交易时间", "收/支", "交易单号", "金额(元)"}.issubset(header_set):
            return index, headers, "wechat"
        if {"交易时间", "收/支", "交易订单号", "金额"}.issubset(header_set):
            return index, headers, "alipay"
    raise ValueError("未识别到微信或支付宝官方账单表头")


def _row_dict(headers: list[str], row: Iterable[Any]) -> dict[str, Any]:
    values = list(row)
    return {
        header: values[index] if index < len(values) else None
        for index, header in enumerate(headers)
        if header
    }


def _parse_datetime(value: Any) -> tuple[str, str]:
    if isinstance(value, datetime):
        return value.date().isoformat(), value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat(), value.isoformat()
    text = _clean(value)
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, pattern)
            return parsed.date().isoformat(), text
        except ValueError:
            continue
    raise ValueError("账单中包含无法识别的交易时间")


def _parse_amount(value: Any) -> float:
    text = _clean(value).replace(",", "").replace("¥", "").replace("￥", "")
    amount = float(text)
    if not math.isfinite(amount) or amount <= 0:
        raise ValueError("账单中包含无效金额")
    return round(amount, 2)


def _direction(value: Any) -> str | None:
    text = _clean(value)
    if text == "收入":
        return "income"
    if text == "支出":
        return "expense"
    return None


def _category(provider: str, trade_category: str, text: str) -> str:
    haystack = f"{trade_category} {text}"
    if provider == "alipay":
        for category, keywords in ALIPAY_CATEGORY_RULES:
            if any(keyword in haystack for keyword in keywords):
                return category
    for category, keywords in GENERAL_CATEGORY_RULES:
        if any(keyword in haystack for keyword in keywords):
            return category
    return "其他"


def _note(counterparty: str, product: str, trade_category: str) -> str:
    values = [value for value in (counterparty, product) if value and value != "/"]
    if len(values) > 1 and values[1] in values[0]:
        values = values[:1]
    return (" · ".join(values) or trade_category or "账单导入")[:80]


def _fingerprint(
    provider: str,
    external_id: str,
    occurred_at: str,
    direction: str,
    amount: float,
    counterparty: str,
    product: str,
) -> str:
    identity = external_id if external_id and external_id != "/" else "|".join(
        (occurred_at, direction, f"{amount:.2f}", counterparty, product)
    )
    return hashlib.sha256(f"{provider}|{identity}".encode()).hexdigest()


def _user_import_key(user_id: str, fingerprint: str) -> str:
    return hashlib.sha256(f"{user_id}|{fingerprint}".encode()).hexdigest()


def parse_bill(filename: str, content: bytes) -> dict[str, Any]:
    if not content:
        raise ValueError("账单文件为空")
    if len(content) > MAX_IMPORT_FILE_BYTES:
        raise ValueError("账单文件不能超过 12MB")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "xlsx":
        rows = _xlsx_rows(content)
    elif suffix == "csv":
        rows = _csv_rows(content)
    else:
        raise ValueError("仅支持微信 XLSX 或支付宝 CSV 账单")

    header_index, headers, provider = _find_header(rows)
    source = f"{provider}-import"
    transactions: list[dict[str, Any]] = []
    skipped = 0
    total_rows = 0

    for raw_row in rows[header_index + 1 :]:
        if not any(_clean(value) for value in raw_row):
            continue
        total_rows += 1
        if total_rows > MAX_IMPORT_ROWS:
            raise ValueError(f"单次最多导入 {MAX_IMPORT_ROWS} 笔交易")
        item = _row_dict(headers, raw_row)
        transaction_type = _direction(item.get("收/支"))
        status = _clean(item.get("当前状态") or item.get("交易状态"))
        if not transaction_type or any(marker in status for marker in FAILED_STATUS_MARKERS):
            skipped += 1
            continue
        try:
            transaction_date, occurred_at = _parse_datetime(item.get("交易时间"))
            amount = _parse_amount(item.get("金额(元)") or item.get("金额"))
        except (TypeError, ValueError):
            skipped += 1
            continue

        trade_category = _clean(item.get("交易类型") or item.get("交易分类"))
        counterparty = _clean(item.get("交易对方"))
        product = _clean(item.get("商品") or item.get("商品说明"))
        external_id = _clean(item.get("交易单号") or item.get("交易订单号"))
        fingerprint = _fingerprint(
            provider,
            external_id,
            occurred_at,
            transaction_type,
            amount,
            counterparty,
            product,
        )
        text = f"{counterparty} {product} {trade_category}"
        transactions.append(
            {
                "fingerprint": fingerprint,
                "externalId": external_id,
                "amount": amount,
                "type": transaction_type,
                "category": _category(provider, trade_category, text),
                "note": _note(counterparty, product, trade_category),
                "date": transaction_date,
                "occurredAt": occurred_at,
                "source": source,
                "status": status,
                "paymentMethod": _clean(item.get("支付方式") or item.get("收/付款方式")),
            }
        )

    if not transactions:
        raise ValueError("账单中没有可导入的收入或支出记录")
    return {
        "provider": provider,
        "providerLabel": "微信支付" if provider == "wechat" else "支付宝",
        "fileName": filename,
        "totalRows": total_rows,
        "skippedCount": skipped,
        "transactions": transactions,
    }


def preview_bill(db: Session, user_id: str, parsed: dict[str, Any]) -> dict[str, Any]:
    transactions = parsed["transactions"]
    keys = [_user_import_key(user_id, item["fingerprint"]) for item in transactions]
    existing = set(
        db.scalars(
            select(Transaction.import_key)
            .join(Ledger, Transaction.ledger_id == Ledger.id)
            .where(Ledger.user_id == user_id, Transaction.import_key.in_(keys))
        )
    )
    preview_rows = []
    duplicate_count = 0
    for item, import_key in zip(transactions, keys, strict=True):
        duplicate = import_key in existing
        duplicate_count += int(duplicate)
        row = {key: value for key, value in item.items() if key != "fingerprint"}
        row.update({"importKey": import_key, "duplicate": duplicate})
        preview_rows.append(row)
    return {
        **{key: value for key, value in parsed.items() if key != "transactions"},
        "importableCount": len(preview_rows) - duplicate_count,
        "duplicateCount": duplicate_count,
        "transactions": preview_rows,
    }


def import_bill_transactions(
    db: Session,
    user_id: str,
    ledger_id: str,
    transactions: list[dict[str, Any]],
) -> tuple[list[Transaction], int]:
    ledger = db.scalar(select(Ledger).where(Ledger.id == ledger_id, Ledger.user_id == user_id))
    if not ledger:
        raise ValueError("目标账本不存在")
    if not transactions:
        raise ValueError("请至少选择一笔交易")
    if len(transactions) > MAX_IMPORT_ROWS:
        raise ValueError(f"单次最多导入 {MAX_IMPORT_ROWS} 笔交易")

    requested_keys: set[str] = set()
    for item in transactions:
        import_key = str(item.get("importKey", ""))
        if not re.fullmatch(r"[0-9a-f]{64}", import_key):
            raise ValueError("导入交易标识无效，请重新预览账单")
        requested_keys.add(import_key)
    existing = set(
        db.scalars(select(Transaction.import_key).where(Transaction.import_key.in_(requested_keys)))
    )
    added: list[Transaction] = []
    seen = set(existing)
    skipped = 0

    for item in transactions:
        import_key = str(item.get("importKey", ""))
        if import_key in seen:
            skipped += 1
            continue
        source = str(item.get("source", ""))
        transaction_type = str(item.get("type", ""))
        try:
            amount = float(item.get("amount", 0))
        except (TypeError, ValueError) as exc:
            raise ValueError("导入交易金额格式不正确") from exc
        if source not in IMPORT_SOURCES or transaction_type not in {"income", "expense"}:
            raise ValueError("导入交易格式不正确")
        if not math.isfinite(amount) or amount <= 0 or amount > 999_999_999_999.99:
            raise ValueError("导入交易金额超出允许范围")
        try:
            transaction_date = datetime.strptime(str(item.get("date")), "%Y-%m-%d").date()
        except ValueError as exc:
            raise ValueError("导入交易日期格式不正确") from exc
        transaction = Transaction(
            ledger_id=ledger.id,
            amount=Decimal(str(round(amount, 2))),
            type=transaction_type,
            category=(str(item.get("category") or "其他").strip() or "其他")[:24],
            note=(str(item.get("note") or "账单导入").strip() or "账单导入")[:80],
            transaction_date=transaction_date,
            source=source,
            import_key=import_key,
        )
        db.add(transaction)
        added.append(transaction)
        seen.add(import_key)

    db.commit()
    return added, skipped
